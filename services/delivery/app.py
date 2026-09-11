"""HTTP-only orchestration and durable local delivery. No planning logic lives here."""
from __future__ import annotations

from collections import Counter
from contextlib import asynccontextmanager, contextmanager
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import uuid

import httpx
from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shared.contracts import LayoutResult, UnderstandingPackage

ROOT = Path(__file__).resolve().parents[2]
STATIC = Path(__file__).parent / 'static'
RUNTIME = ROOT / 'runtime' / 'delivery'
DB = RUNTIME / 'jobs.sqlite3'
UNDERSTANDING_URL = os.environ.get('L3_UNDERSTANDING_URL', 'http://127.0.0.1:8101')
PLANNING_URL = os.environ.get('L3_PLANNING_URL', 'http://127.0.0.1:8102')
MAX_UPLOAD = 25 * 1024 * 1024


@contextmanager
def connect():
    connection = sqlite3.connect(DB, timeout=30)
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def initialize():
    RUNTIME.mkdir(parents=True, exist_ok=True)
    with connect() as db:
        db.execute('CREATE TABLE IF NOT EXISTS jobs (id TEXT PRIMARY KEY, status TEXT NOT NULL, progress INTEGER NOT NULL, message TEXT NOT NULL, error TEXT, result TEXT)')
        db.execute("UPDATE jobs SET status='failed', message='服务重启导致处理停止，请重新上传。', error=? WHERE status NOT IN ('completed', 'failed')", (json.dumps({'code': 'PROCESS_INTERRUPTED', 'message': '服务重启，请重新上传 CAD。'}, ensure_ascii=False),))


def update(job_id, status, progress, message, error=None, result=None):
    with connect() as db:
        db.execute('UPDATE jobs SET status=?, progress=?, message=?, error=?, result=? WHERE id=?', (status, progress, message, json.dumps(error, ensure_ascii=False) if error else None, json.dumps(result, ensure_ascii=False, allow_nan=False) if result else None, job_id))


def get_job(job_id):
    with connect() as db:
        db.row_factory = sqlite3.Row
        row = db.execute('SELECT * FROM jobs WHERE id=?', (job_id,)).fetchone()
    if row is None:
        raise HTTPException(404, detail={'code': 'JOB_NOT_FOUND', 'message': '未找到此任务。'})
    item = dict(row)
    item['error'] = json.loads(item['error']) if item['error'] else None
    item['result'] = json.loads(item['result']) if item['result'] else None
    return item


def validate_delivery(result):
    parsed = LayoutResult.model_validate(result)
    if parsed.validation.get('status') != 'PASS':
        raise ValueError('布局几何安全验证未通过，不能发布结果。')
    if len({s.id for s in parsed.shelves}) != len(parsed.shelves):
        raise ValueError('货架实例编号重复。')
    counts = Counter(s.material_id for s in parsed.shelves)
    if len({b.material_id for b in parsed.bom}) != len(parsed.bom):
        raise ValueError('物料清单包含重复条目。')
    if dict(counts) != {b.material_id: b.quantity for b in parsed.bom}:
        raise ValueError('货架实例与物料清单数量不一致。')
    specs = {b.material_id: (b.length_mm, b.depth_mm, b.default_level_count) for b in parsed.bom}
    if any(specs[s.material_id] != (s.length_mm, s.depth_mm, s.default_level_count) for s in parsed.shelves):
        raise ValueError('货架实例与物料清单规格不一致。')
    return parsed.model_dump(mode='json')


def make_workbook(result):
    result = validate_delivery(result)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '设计级货架清单'
    sheet.append(['物料编号', '长度 (mm)', '深度 (mm)', '默认层数', '设计使用数量', '高度 (mm)'])
    for row in result['bom']:
        sheet.append([row['material_id'], row['length_mm'], row['depth_mm'], row['default_level_count'], row['quantity'], '缺失，需确认'])
    sheet.append(['合计', None, None, None, len(result['shelves'])])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:F{len(result["bom"]) + 1}'
    for col, width in zip('ABCDEF', [30, 18, 18, 16, 22, 23]):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='173D39')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='center')
    notes = workbook.create_sheet('说明与来源')
    for row in [
        ['清单性质', '设计级货架物料清单；不代表供应商报价、库存、采购可得性或采购真实性。'],
        ['数量含义', '本布局中使用的完整货架实例数；层板不是独立货架数量。'],
        ['高度缺失', '原始物料未提供高度；3D 高度仅为显示假设，不作为采购规格。'],
        ['默认层数', '来自获批物料配置；采购前需确认层数、结构、载荷和高度。'],
        ['商品尺寸', '未从商品名称或规格推断包装尺寸，未进行 SKU 上架或容量优化。'],
        ['源 CAD', result['source']['filename']],
        ['源 SHA256', result['source']['sha256']],
        ['基础过道 (mm)', result['rules'].get('aisle_mm')],
        ['几何验证', result['validation']['status']],
    ]:
        notes.append(row)
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 100
    for row in notes:
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        notes.row_dimensions[row[0].row].height = 34
    # Uploaded source names are text, never spreadsheet formulae.
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def validate_provenance(result, package):
    """Prevent a valid but unrelated service result from being published."""
    expected = package.model_dump(mode='json')
    if result['source'] != expected['source'] or result['spatial'] != expected['spatial']:
        raise ValueError('规划结果的来源或空间数据与当前上传任务不一致。')
    templates = {row.material_id: row for row in package.business.templates}
    for row in result['bom']:
        template = templates.get(row['material_id'])
        if template is None or (row['length_mm'], row['depth_mm'], row['default_level_count']) != (template.length_mm, template.depth_mm, template.default_level_count):
            raise ValueError('规划结果使用了未批准或不匹配的物料规格。')


def checked_json(response):
    if response.is_error:
        try:
            detail = response.json().get('detail', {})
        except ValueError:
            detail = {}
        if not isinstance(detail, dict):
            detail = {}
        raise UpstreamError(detail.get('code', 'UPSTREAM_REJECTED'), detail.get('message', '服务拒绝此输入。'))
    return response.json()


class UpstreamError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def run_job(job_id, source_path, filename):
    try:
        update(job_id, 'understanding', 15, '正在理解 CAD，并加载内部商品与物料配置…')
        with httpx.Client(timeout=300, trust_env=False) as client:
            with open(source_path, 'rb') as source:
                understood = checked_json(client.post(UNDERSTANDING_URL + '/understand', files={'file': (filename, source, 'application/octet-stream')}))
            package = UnderstandingPackage.model_validate(understood)
            update(job_id, 'planning', 55, '正在生成候选布局、检查几何安全与基础过道…')
            result = checked_json(client.post(PLANNING_URL + '/plan', json=package.model_dump(mode='json')))
        result = validate_delivery(result)
        validate_provenance(result, package)
        folder = source_path.parent
        (folder / 'layout.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
        (folder / 'materials.xlsx').write_bytes(make_workbook(result))
        update(job_id, 'completed', 100, '规划完成，几何检查通过。', result=result)
    except UpstreamError as exc:
        update(job_id, 'failed', 0, str(exc), error={'code': exc.code, 'message': str(exc)})
    except httpx.RequestError:
        message = '处理服务暂时无法连接或响应超时，请检查三个服务均已启动后重试。'
        update(job_id, 'failed', 0, message, error={'code': 'SERVICE_UNAVAILABLE', 'message': message})
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else '处理未完成，请查看服务日志并重新上传。'
        import logging
        logging.exception('Delivery job %s failed', job_id)
        update(job_id, 'failed', 0, message, error={'code': 'DELIVERY_VALIDATION_FAILED', 'message': message})


@asynccontextmanager
async def lifespan(app):
    initialize()
    yield


app = FastAPI(title='L3 门店规划 · 网页交付', lifespan=lifespan)
app.mount('/static', StaticFiles(directory=STATIC), name='static')


@app.get('/health')
def health():
    return {'status': 'ok', 'service': 'delivery'}


@app.get('/')
def home():
    return FileResponse(STATIC / 'index.html')


@app.post('/api/jobs', status_code=202)
async def submit(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    filename = (file.filename or '').replace('\\', '/').rsplit('/', 1)[-1]
    if not filename.lower().endswith('.dxf'):
        raise HTTPException(422, detail={'code': 'DXF_REQUIRED', 'message': '请上传一份毫米单位的 DXF CAD；DWG 请先转换为 DXF。'})
    job_id = uuid.uuid4().hex
    folder = RUNTIME / job_id
    folder.mkdir(parents=True)
    source = folder / 'source.dxf'
    size = 0
    try:
        with source.open('wb') as target:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_UPLOAD:
                    raise HTTPException(413, detail={'code': 'FILE_TOO_LARGE', 'message': '文件超过 25 MB 限制。'})
                target.write(chunk)
        if size == 0:
            raise HTTPException(422, detail={'code': 'EMPTY_FILE', 'message': 'CAD 文件为空。'})
    except Exception:
        source.unlink(missing_ok=True)
        folder.rmdir()
        raise
    finally:
        await file.close()
    with connect() as db:
        db.execute('INSERT INTO jobs (id,status,progress,message) VALUES (?,?,?,?)', (job_id, 'queued', 5, '已接收 CAD，等待处理…'))
    background_tasks.add_task(run_job, job_id, source, filename)
    return {'id': job_id, 'status': 'queued'}


@app.get('/api/jobs/{job_id}')
def status(job_id: str):
    return get_job(job_id)


def artifact(job_id, name, media_type):
    if get_job(job_id)['status'] != 'completed':
        raise HTTPException(409, detail={'code': 'RESULT_NOT_READY', 'message': '任务尚未成功完成，暂不能下载。'})
    return FileResponse(RUNTIME / job_id / name, media_type=media_type, filename=name)


@app.get('/api/jobs/{job_id}/layout.json')
def layout(job_id: str):
    return artifact(job_id, 'layout.json', 'application/json')


@app.get('/api/jobs/{job_id}/materials.xlsx')
def materials(job_id: str):
    return artifact(job_id, 'materials.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# Preserve the rejected shell solely for explicit historical regression tests.
# The standard three-service launcher now exposes only the gated v2 workbench.
failed_baseline_app = app
from services.delivery.workbench import app
