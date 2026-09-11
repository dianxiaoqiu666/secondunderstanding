"""Design-level workbook from validated final module instances, using openpyxl."""
from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from shared.planning_v2 import LayoutV2
from shared.operations_planning import OperationsLayout


def make_workbook(value):
    operational=isinstance(value,OperationsLayout) or (isinstance(value,dict) and 'operational_schema_version' in value)
    model=OperationsLayout if operational else LayoutV2
    result=model.model_validate(value.model_dump(mode='json') if operational and isinstance(value,OperationsLayout) else value)
    if result.design_status!='VALIDATED' or result.validation.get('status')!='PASS':
        raise ValueError('Only validated final layouts may be exported')
    if operational and (result.route_evaluation.get('status')!='PASS'
            or result.validation.get('route_evidence',{}).get('status')!='PASS'):
        raise ValueError('Operational layouts require complete validated route evidence')
    counts=Counter(s.material_id for s in result.shelves)
    if counts!={row.material_id:row.quantity for row in result.bom}:
        raise ValueError('Module and material counts differ')
    workbook=Workbook()
    sheet=workbook.active
    sheet.title='设计级货架清单'
    sheet.append(['物料编号','长度 (mm)','深度 (mm)','默认层数','使用数量','总层数','高度 (mm)'])
    for row in result.bom:
        if row.total_level_count!=row.quantity*row.default_level_count:
            raise ValueError('Default level total differs')
        sheet.append([row.material_id,row.length_mm,row.depth_mm,row.default_level_count,
                      row.quantity,row.total_level_count,'原始配置未提供'])
    sheet.append(['合计',None,None,None,len(result.shelves),sum(r.total_level_count for r in result.bom)])
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=f'A1:G{len(result.bom)+1}'
    for column,width in zip('ABCDEFG',[28,18,18,16,16,18,26]):
        sheet.column_dimensions[column].width=width
    for cell in sheet[1]:
        cell.fill=PatternFill('solid',fgColor='17665E')
        cell.font=Font(bold=True,color='FFFFFF')
    instances=workbook.create_sheet('货架实例')
    instances.append(['货架编号','货架排','物料编号','X (mm)','Y (mm)','方向 (deg)','默认层数','品类'])
    category_by_shelf={sid:row['category'] for row in result.products.category_assignments for sid in row['shelf_ids']}
    for run in result.runs:
        for shelf in run.modules:
            instances.append([shelf.id,run.run_id,shelf.material_id,shelf.x_mm,shelf.y_mm,
                              shelf.rotation_deg,shelf.default_level_count,category_by_shelf.get(shelf.id,'未分配')])
    instances.freeze_panes='A2'
    instances.auto_filter.ref=instances.dimensions
    for column in 'ABCDEFGH':instances.column_dimensions[column].width=22
    notes=workbook.create_sheet('说明与来源')
    for row in [
        ['清单性质','设计级完整货架模块清单，不是零件采购BOM、报价、库存或供货承诺。'],
        ['默认层数','来自内部获批物料配置。未提供货架高度、层间净高、承重与完整零件结构。'],
        ['显示假设','作业布局为二维明确场景；货架高度、净层板尺寸与结构零件尚未提供。' if operational else '三维货架高1800mm、墙高2600mm/厚100mm仅用于显示，不作为采购规格或商品净高。'],
        ['商品规划','按来源品类分配架位；真实SKU包装尺寸缺失，不生成精确上架、容量或销量结论。'],
        ['明确空间输入' if operational else '源CAD',result.source.filename],['源SHA256',result.source.sha256],
        ['空间来源',result.space.confirmation.state],['基础排间过道 (mm)',result.rules.aisle_width_mm],
        ['排端通道 (mm)',result.rules.end_aisle_mm],['几何检查',result.validation['status']],
        ['算法范围','全部取货面及固定任务经过当前净宽路径检查；不证明真实门店效果或订单全局最短。' if operational else '确定性连续排启发式；局部几何和过道检查不等于全店可达或消防认证。'],
        ['结果SHA256',result.metrics.get('deterministic_sha256','')],
    ]:notes.append(row)
    if operational:
        notes.append(['员工作业评测','明确测试位置及固定合成任务；人工产品验收仍待复验，不代表真实门店效率。'])
        notes.append(['双面物料口径','当前为两套独立单面模块背靠背；两侧分别计完整模块，组合数另列，未提供共用结构零件。'])
        notes.append(['空间指标','长×单侧规划深度×默认层数仅为名义层板面积；净层板尺寸与商品容量 UNKNOWN。'])
        notes.append(['作业输入SHA256',result.validation.get('operations_input_sha256','')])
        assembly_sheet=workbook.create_sheet('单双面组合')
        assembly_sheet.append(['组合编号','结构测试类型','单侧排编号','独立模块数','配对双面模块组数','各单侧深度(mm)','结构间隙(mm)','组合总深(mm)'])
        by_run={run.run_id:run for run in result.runs}
        for item in result.assemblies:
            assembly_sheet.append([item.id,item.kind,', '.join(item.run_ids),
                sum(len(by_run[rid].modules) for rid in item.run_ids),len(item.bay_pairs),
                ' + '.join(str(depth) for depth in item.side_depths_mm),item.structure_gap_mm,item.total_depth_mm])
        assembly_sheet.freeze_panes='A2'
        for column in 'ABCDEFGH':assembly_sheet.column_dimensions[column].width=26
    notes.column_dimensions['A'].width=28
    notes.column_dimensions['B'].width=100
    for row in notes:
        notes.row_dimensions[row[0].row].height=34
        row[1].alignment=Alignment(wrap_text=True,vertical='top')
    # External names and categories are always literal text, including '=...'.
    for tab in workbook:
        for row in tab:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
    output=BytesIO()
    workbook.save(output)
    return output.getvalue()
