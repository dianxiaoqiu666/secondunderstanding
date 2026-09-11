"""An observed-object workbook, explicitly distinct from a new design BOM."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from shared.existing_design import verify_existing_digest


def make_existing_workbook(value):
    verify_existing_digest(value)
    book=Workbook();sheet=book.active;sheet.title='原图已识别货架'
    sheet.append(['物料匹配','实测长度 mm','实测深度 mm','配置默认层数','原图数量','性质'])
    for row in value['bom']:
        sheet.append([row.get('material_id') or '未匹配',row['length_mm'],row['depth_mm'],
                      row.get('default_level_count') or '未知',row['quantity'],'原图观察，不是新方案采购清单'])
    sheet.append(['合计',None,None,None,len(value['shelves'])])
    instances=book.create_sheet('原始位置')
    instances.append(['对象','X mm','Y mm','方向 deg','源实体','真实高度'])
    for shelf in value['shelves']:
        instances.append([shelf['id'],shelf['x_mm'],shelf['y_mm'],shelf['rotation_deg'],
                          ', '.join(shelf['source_handles']),shelf.get('measured_height_mm') or '未知'])
    excluded=book.create_sheet('未计入对象')
    excluded.append(['对象','分类','原因'])
    for group,label in [('legend_samples','图例或规格样本'),('unknown_objects','无法确定')]:
        for item in value[group]:excluded.append([item.get('id',item.get('handle','')),label,item.get('reason','')])
    notes=book.create_sheet('用途及限制')
    for row in [['用途','查看上传图中的已有设计，没有自动重新摆放，没有评估新规划几何安全。'],
                ['来源',value['source']['filename']],['SHA256',value['source']['sha256']],
                ['原分类',value['source_classification']],['本次用途','existing_design'],
                ['展示参数','墙/货架展示高度和未知层数的展示层板不是实测规格。'],
                ['未呈现', '；'.join(value['summary'].get('omissions',[]))]]:notes.append(row)
    for tab in book:
        tab.freeze_panes='A2'
        for cell in tab[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17665E')
        for col in 'ABCDEF':tab.column_dimensions[col].width=30
        for row in tab:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
    notes.column_dimensions['B'].width=110
    out=BytesIO();book.save(out);return out.getvalue()
