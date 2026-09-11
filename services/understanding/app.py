"""Offline, fail-closed CAD understanding; no reference-design input."""
from __future__ import annotations

import hashlib
import io
import json
import math
from pathlib import Path

import ezdxf
from ezdxf.lldxf.tagger import ascii_tags_loader
from ezdxf.lldxf.tags import group_tags
from ezdxf.lldxf.types import DOUBLE
from fastapi import FastAPI, File, HTTPException, UploadFile
from openpyxl import load_workbook
from shapely import make_valid, normalize
from shapely.geometry import Polygon

from shared.contracts import Business, Exclusion, PolygonData, Source, Spatial, Template, UnderstandingPackage, Wall

ROOT = Path(__file__).resolve().parents[2]
LEGACY_SHA = 'bd92b40c74cbbe75001b470f20afbaa3ca4aa3da0d2cb425d98d8c1d015defd0'
PRODUCT_PATH = 'input/production/products/副本门店商品导出_1000370158187220260814.xlsx'
PRODUCT_SHA = 'e66bdeed9321ce29e569267a059bcb6d0f7f17cc53af1037b0f61da5c24cb169'
MATERIAL_PATH = 'input/production/materials/material_templates.v1.json'
MATERIAL_SHA = 'b63863780526fefcdb35f9cde07673ed538b287251da77b80e8b69612a4cb024'
SCOPE_LAYERS = {'PLANNING_SCOPE', 'HUMAN_PLANNING_SCOPE'}
WALL_LAYERS = {'墙体', 'WALL', 'WALLS'}
HOLE_LAYERS = {'PLANNING_HOLE', 'PLANNING_HOLES'}
EXCLUSION_LAYERS = {'EXCLUSION', 'EXCLUSION_ZONE', 'NO_PLACEMENT'}


class InputError(ValueError):
    def __init__(self, code: str, message: str):
        self.code, self.message = code, message
        super().__init__(message)


def fail(code, message):
    raise InputError(code, message)


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def point(value):
    coords = tuple(float(x) for x in value)
    if not all(math.isfinite(x) for x in coords) or (len(coords) > 2 and abs(coords[2]) > 1e-9):
        fail('NON_PLANAR_GEOMETRY', '仅支持有限坐标、Z=0 的二维毫米 CAD。')
    return coords[:2]


def check_plane(entity):
    extrusion = tuple(entity.dxf.get('extrusion', (0, 0, 1)))
    if extrusion != (0, 0, 1):
        fail('NON_PLANAR_GEOMETRY', f'实体 {entity.dxf.handle} 的法向量不是二维正向 Z。')
    # Curves such as ARC encode elevation in their center, not an elevation
    # attribute. Their transformed coordinates are checked by point().
    elevation = entity.dxf.get('elevation', 0) if entity.dxf.is_supported('elevation') else 0
    if isinstance(elevation, (int, float)):
        if not math.isfinite(elevation) or abs(elevation) > 1e-9:
            fail('NON_PLANAR_GEOMETRY', 'CAD elevation 必须为 0。')
    else:
        point(elevation)


def polygon_data(poly):
    poly = normalize(poly)
    return dict(boundary_mm=list(poly.exterior.coords), holes_mm=[list(r.coords) for r in poly.interiors])


def strict_polyline(entity):
    if entity.dxftype() != 'LWPOLYLINE':
        fail('UNSUPPORTED_BOUNDARY', f'实体 {entity.dxf.handle} 必须为闭合 LWPOLYLINE。')
    check_plane(entity)
    vertices = list(entity.get_points())
    if any(v[4] != 0 for v in vertices):
        fail('CURVED_BOUNDARY_UNSUPPORTED', '当前输入契约要求范围、孔洞和禁放区域使用直线段。')
    if not entity.closed:
        fail('SCOPE_NOT_CLOSED', f'实体 {entity.dxf.handle} 未设置闭合标志；请在 CAD 中闭合。')
    coords = [point(v[:2]) for v in vertices]
    if len(coords) < 3:
        fail('INVALID_BOUNDARY', '边界至少需要三个不同顶点。')
    poly = Polygon(coords)
    if not poly.is_valid or not math.isfinite(poly.area) or poly.area <= 0:
        fail('INVALID_BOUNDARY', f'实体 {entity.dxf.handle} 自交、退化或无效。')
    return poly


def hatch_polygons(entity):
    """Interpret NORMAL style using valid, non-crossing loop containment parity."""
    check_plane(entity)
    if entity.dxf.hatch_style != 0:
        fail('HATCH_STYLE_UNSUPPORTED', f'HATCH {entity.dxf.handle} 仅支持 style 0。')
    loops = []
    for path in entity.paths:
        if type(path).__name__ == 'PolylinePath':
            if not path.is_closed or any(v[2] != 0 for v in path.vertices):
                fail('HATCH_PATH_UNSUPPORTED', f'HATCH {entity.dxf.handle} 含开放或曲线路径。')
            coords = [point(v[:2]) for v in path.vertices]
        elif type(path).__name__ == 'EdgePath':
            if not path.edges or any(type(e).__name__ != 'LineEdge' for e in path.edges):
                fail('HATCH_PATH_UNSUPPORTED', f'HATCH {entity.dxf.handle} 仅支持连续直线边。')
            for i, edge in enumerate(path.edges):
                if point(edge.end) != point(path.edges[(i + 1) % len(path.edges)].start):
                    fail('HATCH_PATH_OPEN', f'HATCH {entity.dxf.handle} 边界不连续闭合。')
            coords = [point(e.start) for e in path.edges]
        else:
            fail('HATCH_PATH_UNSUPPORTED', f'HATCH {entity.dxf.handle} 路径不受支持。')
        if len(coords) < 3:
            fail('HATCH_INVALID', 'HATCH 边界顶点不足。')
        poly = Polygon(coords)
        if not poly.is_valid or not math.isfinite(poly.area) or poly.area <= 0:
            fail('HATCH_INVALID', f'HATCH {entity.dxf.handle} 含无效边界，不能部分使用。')
        for previous in loops:
            if previous.boundary.intersects(poly.boundary) or (previous.intersects(poly) and not (previous.contains(poly) or poly.contains(previous))):
                fail('HATCH_LOOPS_AMBIGUOUS', f'HATCH {entity.dxf.handle} 含相交或接触边界。')
        loops.append(poly)
    if not loops:
        fail('HATCH_INVALID', 'HATCH 不含边界。')
    result = loops[0]
    for poly in loops[1:]:
        result = result.symmetric_difference(poly)
    if result.geom_type not in {'Polygon', 'MultiPolygon'} or not result.is_valid:
        fail('HATCH_INVALID', 'HATCH 无法解释为有效排除区域。')
    polys = [result] if result.geom_type == 'Polygon' else list(result.geoms)
    return sorted(polys, key=lambda p: (p.bounds, p.area))


def validate_raw_tags(text: str):
    """Check uncorrected ezdxf tags before entity loading applies defaults/fixes."""
    for group in group_tags(ascii_tags_loader(io.StringIO(text, newline=None))):
        for tag in group:
            if tag.code in DOUBLE and not math.isfinite(float(tag.value)):
                fail('NONFINITE_DXF_VALUE', '原始 DXF 含 NaN/Infinity 坐标或数值，请修复图纸。')
        if group[0].value == 'HATCH':
            styles = [tag.value.strip() for tag in group if tag.code == 75]
            if len(styles) != 1 or styles[0] != '0':
                fail('HATCH_STYLE_UNSUPPORTED', '原始 HATCH 必须明确且唯一指定 style 0；缺失、非法或其他样式均拒绝。')


def parse_cad(data: bytes) -> Spatial:
    if not data or len(data) > 25 * 1024 * 1024:
        fail('CAD_SIZE_INVALID', '请上传非空且不超过 25 MB 的 DXF。')
    try:
        text = data.decode('utf-8-sig')
        validate_raw_tags(text)
        doc = ezdxf.read(io.StringIO(text, newline=None))
    except InputError:
        raise
    except Exception as exc:
        fail('CAD_PARSE_FAILED', f'无法读取 DXF；请导出 UTF-8 DXF 2007 或更新格式。({type(exc).__name__})')
    if doc.units != 4:
        fail('UNITS_NOT_MM', 'CAD 必须明确设置 INSUNITS=4（毫米）。')
    audit = doc.audit()
    if audit.has_errors:
        fail('CAD_AUDIT_FAILED', 'DXF 结构审计失败，请修复图纸。')
    # The exact approved source contains known plotting/annotation repairs only.
    if audit.has_fixes and digest(data) != LEGACY_SHA:
        fail('CAD_REPAIR_REQUIRED', 'DXF 需要内部修复，请在 CAD 软件中审计修复后重新导出。')
    entities = list(doc.modelspace())
    scopes = [e for e in entities if e.dxf.layer.upper() in SCOPE_LAYERS]
    provenance = {'parser': 'ezdxf', 'geometry_kernel': 'GEOS/Shapely', 'cad_sha256': digest(data)}
    provenance['dxf_audit_fixes'] = [{'code': int(item.code), 'message': item.message} for item in audit.fixes]
    warnings = []
    if digest(data) == LEGACY_SHA:
        if scopes:
            fail('SCOPE_AMBIGUOUS', '兼容图纸出现额外规划范围。')
        entity = doc.entitydb.get('180')
        coords = [point(v[:2]) for v in entity.get_points()]
        repaired = make_valid(Polygon(coords))
        polys = [g for g in repaired.geoms if g.geom_type == 'Polygon']
        if len(polys) != 1:
            fail('LEGACY_SCOPE_INVALID', '获批兼容范围拓扑不符合预期。')
        scope = polys[0]
        lines = []
        for geom in repaired.geoms:
            if geom.geom_type == 'MultiLineString':
                lines.extend([list(line.coords) for line in geom.geoms])
            elif geom.geom_type == 'LineString':
                lines.append(list(geom.coords))
        provenance['scope'] = {'mode': 'approved_sha256_handle_compatibility', 'handle': '180', 'source_closed': False,
            'operations': ['closed_first_last_vertex', 'make_valid'], 'closing_segment_mm': [coords[-1], coords[0]],
            'closing_segment_length_mm': math.dist(coords[-1], coords[0]), 'degenerate_lines_mm': lines,
            'area_mm2': scope.area, 'hole_count': len(scope.interiors)}
        warnings.append('该精确获批 CAD 使用 handle 180 历史兼容解释：补首尾线并 make_valid；保留孔洞及退化线审计。')
    else:
        if len(scopes) != 1:
            fail('SCOPE_COUNT_INVALID', f'必须有且仅有一个 PLANNING_SCOPE/HUMAN_PLANNING_SCOPE，当前 {len(scopes)} 个。')
        scope = strict_polyline(scopes[0])
        provenance['scope'] = {'mode': 'explicit_layer', 'handle': scopes[0].dxf.handle, 'operations': []}
    holes = [strict_polyline(e) for e in entities if e.dxf.layer.upper() in HOLE_LAYERS]
    for index, hole in enumerate(holes):
        if not scope.contains(hole) or scope.boundary.intersects(hole.boundary) or any(hole.intersects(other) for other in holes[:index]):
            fail('HOLE_INVALID', '孔洞必须严格位于规划范围内且互不接触、重叠。')
    if holes:
        scope = Polygon(scope.exterior.coords, [list(r.coords) for r in scope.interiors] + [list(p.exterior.coords) for p in holes])
    walls, exclusions = [], []
    for entity in entities:
        kind, layer, handle = entity.dxftype(), entity.dxf.layer.upper(), entity.dxf.handle
        if kind == 'HATCH':
            for index, poly in enumerate(hatch_polygons(entity)):
                exclusions.append(Exclusion(id=f'hatch-{handle}-{index}', **polygon_data(poly)))
        elif layer in EXCLUSION_LAYERS:
            exclusions.append(Exclusion(id=f'exclusion-{handle}', **polygon_data(strict_polyline(entity))))
        elif layer in WALL_LAYERS:
            if kind == 'LINE':
                start, end = point(entity.dxf.start), point(entity.dxf.end)
                if start == end:
                    fail('WALL_DEGENERATE', f'墙 {handle} 长度为零。')
                walls.append(Wall(id=f'wall-{handle}', start_mm=start, end_mm=end))
            elif kind not in {'TEXT', 'MTEXT', 'DIMENSION'}:
                fail('WALL_GEOMETRY_UNSUPPORTED', f'墙图层实体 {handle} ({kind}) 不受支持；请转换为二维 LINE。')
        elif kind in {'INSERT', 'XREF', '3DSOLID', 'MESH', 'POLYFACE'}:
            fail('UNSUPPORTED_SPATIAL_ENTITY', f'实体 {handle} ({kind}) 可能隐藏空间约束，请展开为支持的二维实体。')
        elif digest(data) != LEGACY_SHA:
            if layer in SCOPE_LAYERS | HOLE_LAYERS:
                continue  # Already strictly validated above.
            if kind in {'TEXT', 'MTEXT', 'DIMENSION'}:
                continue
            if layer == 'DEFPOINTS' and kind in {'LINE', 'LWPOLYLINE'}:
                continue  # Explicit construction-line layer, never a physical obstacle.
            fail('UNCLASSIFIED_SPATIAL_ENTITY', f'实体 {handle} ({kind}) 位于未约定空间图层 {entity.dxf.layer}；请明确归入墙、范围、孔洞或禁放区域。')
    provenance['wall_handles'] = [w.id.removeprefix('wall-') for w in walls]
    provenance['hatch_handles'] = [e.dxf.handle for e in entities if e.dxftype() == 'HATCH']
    provenance['ignored_annotation_count'] = sum(e.dxftype() in {'TEXT', 'MTEXT', 'DIMENSION'} for e in entities)
    provenance['ignored_construction_handles'] = [e.dxf.handle for e in entities if e.dxf.layer.upper() == 'DEFPOINTS' and e.dxftype() in {'LINE', 'LWPOLYLINE'}]
    if digest(data) == LEGACY_SHA:
        provenance['legacy_auxiliary_handles'] = [e.dxf.handle for e in entities if e.dxf.layer == '0' and e.dxf.handle != '180']
    return Spatial(scope=PolygonData(**polygon_data(scope)), walls=walls, exclusion_zones=exclusions, provenance=provenance, warnings=warnings)


def load_business(root: Path = ROOT) -> Business:
    """Hash both approved inputs before opening either; do not load reference data."""
    try:
        classification = (root / 'DATA_CLASSIFICATION.md').read_text(encoding='utf-8-sig')
        raw = {}
        sources = []
        for relative, expected, role in [(PRODUCT_PATH, PRODUCT_SHA, 'PRODUCTION_INPUT'), (MATERIAL_PATH, MATERIAL_SHA, 'MATERIAL_CONFIG')]:
            if not any(relative in line and expected in line and role in line for line in classification.splitlines()):
                fail('SOURCE_NOT_APPROVED', f'分类清单没有批准 {relative}。')
            data = (root / relative).read_bytes()
            if digest(data) != expected:
                fail('BUSINESS_SOURCE_HASH_MISMATCH', f'内部资料 {relative} 的 SHA256 与批准值不符。')
            raw[relative] = data
            sources.append({'path': relative, 'sha256': expected, 'classification': role})
        catalog = json.loads(raw[MATERIAL_PATH].decode('utf-8-sig'))
        templates = [Template.model_validate(t) for t in catalog['templates']]
        if len({t.material_id for t in templates}) != len(templates):
            fail('MATERIAL_CONFIG_INVALID', '物料 ID 重复。')
        workbook = load_workbook(io.BytesIO(raw[PRODUCT_PATH]), read_only=True, data_only=True)
        try:
            sheet = workbook['门店商品']
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            fields = {name: header.index(name) for name in ['商品条码', '商品名称', '规格名称', '门店店内分类']}
            products = []
            for number, row in enumerate(rows, 2):
                if not any(value is not None for value in row):
                    continue
                products.append({'source_row': number, 'barcode': str(row[fields['商品条码']] or ''),
                    'name': str(row[fields['商品名称']] or ''), 'specification': str(row[fields['规格名称']] or ''),
                    'category': str(row[fields['门店店内分类']] or ''), 'width_mm': None, 'depth_mm': None, 'height_mm': None})
        finally:
            workbook.close()
        if not products:
            fail('PRODUCT_DATA_EMPTY', '商品主数据为空。')
        return Business(products=products, product_count=len(products), templates=templates, sources=sources)
    except InputError:
        raise
    except Exception as exc:
        fail('BUSINESS_SOURCE_INVALID', f'内部商品或物料资料缺失/损坏。({type(exc).__name__})')


def understand_bytes(data: bytes, filename: str = 'store.dxf') -> UnderstandingPackage:
    if not filename.lower().endswith('.dxf'):
        fail('CAD_FORMAT_UNSUPPORTED', '当前支持 DXF，请先将 DWG 导出为 DXF。')
    try:
        spatial = parse_cad(data)
    except InputError:
        raise
    except Exception as exc:
        fail('CAD_GEOMETRY_INVALID', f'CAD 几何数据损坏或不受支持，请修复后重新上传。({type(exc).__name__})')
    business = load_business()
    return UnderstandingPackage(source=Source(filename=Path(filename.replace('\\', '/')).name, sha256=digest(data)), spatial=spatial, business=business)


app = FastAPI(title='L3 Understanding', version='1.0')


@app.get('/health')
def health():
    return {'status': 'ok', 'service': 'understanding'}


@app.post('/understand', response_model=UnderstandingPackage)
def understand(file: UploadFile = File(...)):
    try:
        return understand_bytes(file.file.read(25 * 1024 * 1024 + 1), file.filename or 'store.dxf')
    except InputError as exc:
        raise HTTPException(status_code=422, detail={'code': exc.code, 'message': exc.message}) from exc


# Historical coarse-scope interpretation is not a production HTTP route.
failed_baseline_app = app
app = FastAPI(title='L3 Human-reviewed CAD Understanding', version='2.0')
app.add_api_route('/health',health,methods=['GET'])
from services.understanding.prepare import router
app.include_router(router)
from services.understanding.existing_api import router as existing_router
app.include_router(existing_router)


@app.post('/understand',include_in_schema=False)
def retired_understand():
    raise HTTPException(410,detail={'code':'FAILED_BASELINE_RETIRED','message':'旧粗框自动解释已停用，请先 /prepare 并确认真实空间。'})
