"""Actual ASGI entry contracts; UI acceptance uses separately running services."""
from copy import deepcopy
import json
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient

from shared.explicit_planning import ExplicitPlanningInput, explicit_request
from services.planning.app import app as planning_app
from services.planning.explicit import compare_explicit
from services.planning.engine import PlanningError

ROOT=Path(__file__).resolve().parents[2]


@pytest.fixture
def payload():
    data=json.loads((ROOT/'tools/benchmarks/scenarios_m09.json').read_text(encoding='utf-8-sig'))
    return deepcopy(next(row['input'] for row in data['scenarios'] if row['id']=='entrance'))


def test_http_explicit_input_needs_no_cad_or_signing(payload,monkeypatch):
    from shared import confirmation
    def forbidden(*args,**kwargs):raise AssertionError('Explicit geometry must not need a CAD signing key')
    monkeypatch.setattr(confirmation,'_key',forbidden)
    with TestClient(planning_app) as client:
        response=client.post('/plan-explicit',json=payload)
        again=client.post('/plan-explicit',json=payload)
    assert response.status_code==again.status_code==200
    value=response.json();repeat=again.json()
    assert value['input']==payload
    for phase in ('before','after'):
        layout=value[phase]['layout']
        assert layout==repeat[phase]['layout']
        assert layout['space']['confirmation']['state']=='SYNTHETIC_TEST'
        assert layout['space']['entrances']==payload['space']['reserved_passages']
        assert layout['validation']['status']=='PASS'
        assert sum(row['quantity'] for row in layout['bom'])==len(layout['shelves'])
    assert value['selection']==repeat['selection']


def test_invalid_geometry_and_no_fit_are_not_promoted_to_ready(payload):
    invalid=deepcopy(payload)
    invalid['space']['boundary']['boundary_mm']=[[0,0],[4000,4000],[0,4000],[4000,0],[0,0]]
    tiny=deepcopy(payload)
    tiny['space']['boundary']['boundary_mm']=[[0,0],[1000,0],[1000,1000],[0,1000],[0,0]]
    tiny['space']['reserved_passages']=[]
    with TestClient(planning_app) as client:
        bad=client.post('/plan-explicit',json=invalid)
        small=client.post('/plan-explicit',json=tiny)
    assert bad.status_code==small.status_code==422
    assert small.json()['detail']['code']=='NO_SAFE_LAYOUT'
    assert '不是全局无解' in small.json()['detail']['message']


def test_old_search_failure_does_not_hide_safe_alternative(payload,monkeypatch):
    from services.planning import explicit
    def old_failed(*args,**kwargs):raise PlanningError('VALIDATION_FAILED','old winner failed')
    monkeypatch.setattr(explicit,'generate_layout',old_failed)
    value=explicit.compare_explicit(ExplicitPlanningInput.model_validate(payload))
    assert value['before']['status']=='NOT_FOUND' and value['before']['layout'] is None
    assert value['after']['layout']['validation']['status']=='PASS'


@pytest.fixture
def delivery_client(tmp_path,monkeypatch):
    from services.delivery import workbench
    from services.delivery import algorithm_lab
    monkeypatch.setattr(workbench,'RUNTIME',tmp_path)
    monkeypatch.setattr(workbench,'DB',tmp_path/'jobs.sqlite3')
    calls=[];state={'corrupt':None}
    def transport(request):
        calls.append(str(request.url))
        assert request.url.port==8102 and request.url.path=='/plan-explicit'
        payload=json.loads(request.content)
        value=compare_explicit(ExplicitPlanningInput.model_validate(payload))
        if state['corrupt']=='source':value['input_sha256']='0'*64
        if state['corrupt']=='bom':value['after']['layout']['bom'][0]['quantity']+=1
        return httpx.Response(200,json=value)
    real_client=httpx.Client
    monkeypatch.setattr(algorithm_lab.httpx,'Client',lambda **kwargs:real_client(transport=httpx.MockTransport(transport),**kwargs))
    with TestClient(workbench.app) as client:
        yield client,calls,state


def test_delivery_page_and_only_planning_service_needed(delivery_client):
    client,calls,_=delivery_client
    assert client.get('/algorithm').status_code==200
    catalog=client.get('/api/algorithm/scenarios').json()
    assert catalog['input_kind']=='ALGORITHM_VALIDATION'
    # The original no-door fixtures remain a geometry-only technical endpoint.
    response=client.post('/api/algorithm/geometry/scenarios/column/generate')
    assert response.status_code==200,response.text
    result=response.json()
    assert len(calls)==1 and '/plan-explicit' in calls[0]
    assert result['reference']['status']=='AVAILABLE'
    assert 'density' not in result['reference']['comparison']['reference']
    assert 'density' not in result['reference']['comparison']['automatic']
    assert result['selection']['reference_coordinates_used'] is False
    assert client.post('/api/algorithm/scenarios/unknown/generate').status_code==404


@pytest.mark.parametrize('corrupt',['source','bom'])
def test_delivery_rejects_mismatched_result(delivery_client,corrupt):
    client,_,state=delivery_client;state['corrupt']=corrupt
    response=client.post('/api/algorithm/geometry/scenarios/column/generate')
    assert response.status_code==422
    assert response.json()['detail']['code']=='ALGORITHM_VALIDATION_FAILED'


def test_optimized_modules_accept_existing_product_and_workbook_pipeline(payload):
    # Regression after geometry selection; business records do not choose the space.
    from io import BytesIO
    from openpyxl import load_workbook
    from services.understanding.app import load_business
    from services.planning.search import optimize_layout
    from services.planning.explicit import finalize
    from services.delivery.exports_v2 import make_workbook
    req=explicit_request(ExplicitPlanningInput.model_validate(payload))
    business=load_business()
    assert business.templates==req.business.templates
    req.business=business
    layout,_=optimize_layout(req);layout=finalize(req,layout)
    products=layout.products
    assert sum(row['record_count'] for row in products.category_assignments+products.unassigned_categories)==business.product_count
    assert products.sku_assignments==[]
    assert products.real_sku_status=='UNAVAILABLE-DUE-TO-MISSING-PHYSICAL-DATA'
    assigned=[sid for row in products.category_assignments for sid in row['shelf_ids']]
    assert len(assigned)==len(set(assigned)) and set(assigned)<={s.id for s in layout.shelves}
    book=load_workbook(BytesIO(make_workbook(layout.model_dump(mode='json'))),read_only=True,data_only=True)
    try:
        rows=list(book['设计级货架清单'].values)
        assert rows[-1][4]==len(layout.shelves)
        assert rows[-1][5]==sum(s.default_level_count for s in layout.shelves)
        assert book['货架实例'].max_row-1==len(layout.shelves)
    finally:book.close()
